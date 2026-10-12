from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
NAVY="1F3A5F"; INP=PatternFill("solid",fgColor="FFF4CC"); HDR=PatternFill("solid",fgColor=NAVY); OUT=PatternFill("solid",fgColor="E6F2F1")
thin=Side(style="thin",color="B8C4D0"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
HF=Font(bold=True,color="FFFFFF"); TF=Font(bold=True,size=14,color=NAVY); NF=Font(italic=True,color="5A6B7B")
wb=Workbook()
def hdr(ws,row,vals,col=1):
    for i,v in enumerate(vals):
        c=ws.cell(row=row,column=col+i,value=v); c.fill=HDR; c.font=HF; c.alignment=Alignment(wrap_text=True,vertical="center"); c.border=BOX
def box(ws,rng,fill=None,fmt=None,wrap=False):
    for row in ws[rng]:
        for c in row:
            c.border=BOX
            if fill: c.fill=fill
            if fmt: c.number_format=fmt
            if wrap: c.alignment=Alignment(wrap_text=True,vertical="top")
def widths(ws,ws_w):
    for k,v in ws_w.items(): ws.column_dimensions[k].width=v

# Start
ws=wb.active; ws.title="Start Here"
ws["A1"]="Real Estate Marketing Plan Template"; ws["A1"].font=TF
rows=["Companion to the guide at infina.ai/news/real-estate-marketing-plan/",
"",
"How to use this workbook",
"1. Goal Backward: enter your closings goal and your own conversion rates. The sheet works out how many conversations you need each month and week.",
"2. Budget: enter expected gross commission income and the share you will reinvest. The sheet splits it across channels for your career stage.",
"3. Plan: fill one row per channel with an audience, a weekly action, hours, budget, and the result you will track. The totals check your budget and your hours.",
"4. Quarterly Review: at the end of each quarter, enter spend, conversations, appointments, and closings by source. The sheet calculates spend per closing.",
"",
"Yellow cells are inputs. Teal cells are calculated, so do not type over them.",
"All rates, percentages, and channel shares that come pre-filled are examples. Replace them with your own numbers as soon as you have a few months of data.",
"Nothing here is tax, legal, or compliance advice. Check your state and brokerage rules for advertising before you spend."]
for i,t in enumerate(rows,start=3):
    ws.cell(row=i,column=1,value=t).alignment=Alignment(wrap_text=True,vertical="top")
ws["A5"].font=Font(bold=True,color=NAVY)
ws.column_dimensions["A"].width=110
for r in range(6,14): ws.row_dimensions[r].height=32

# 1 Goal
g=wb.create_sheet("1 Goal Backward")
g["A1"]="Work backward from your goal"; g["A1"].font=TF
hdr(g,3,["Input","Your number","Note"])
data=[("Closings you want this year",10,"Your yearly goal"),
("Share of signed clients who reach closing",0.8,"Example: 8 in 10"),
("Share of appointments that become signed clients",0.5,"Example: 1 in 2"),
("Share of conversations that become appointments",0.2,"Example: 1 in 5")]
for i,(a,b,c) in enumerate(data,start=4):
    g.cell(row=i,column=1,value=a); g.cell(row=i,column=2,value=b); g.cell(row=i,column=3,value=c)
box(g,"A4:C7"); box(g,"B4:B7",INP); g["B4"].number_format="0"
for r in (5,6,7): g.cell(row=r,column=2).number_format="0%"
hdr(g,9,["Result","Number needed","How it is worked out"])
res=[("Signed clients needed","=ROUNDUP(B4/B5,0)","Closings divided by the closing rate"),
("Appointments needed","=ROUNDUP(B10/B6,0)","Signed clients divided by the signing rate"),
("Conversations needed this year","=ROUNDUP(B11/B7,0)","Appointments divided by the appointment rate"),
("Conversations per month","=ROUNDUP(B12/12,0)","Yearly total divided by 12"),
("Conversations per week","=ROUND(B12/52,1)","Yearly total divided by 52")]
for i,(a,b,c) in enumerate(res,start=10):
    g.cell(row=i,column=1,value=a); g.cell(row=i,column=2,value=b); g.cell(row=i,column=3,value=c)
box(g,"A10:C14"); box(g,"B10:B14",OUT)
g["A16"]="If the weekly number looks out of reach for your hours, lower the goal now instead of in December."; g["A16"].font=NF
widths(g,{"A":52,"B":18,"C":46})

# 2 Budget
b=wb.create_sheet("2 Budget")
b["A1"]="Set a budget you can defend"; b["A1"].font=TF
hdr(b,3,["Input","Your number","Note"])
bi=[("Expected gross commission income (GCI) this year",60000,"Example amount"),
("Share of GCI you will reinvest in marketing",0.10,"Example percentage, not a benchmark"),
("Career stage","New agent","Choose New agent or Growing agent")]
for i,(a,v,c) in enumerate(bi,start=4):
    b.cell(row=i,column=1,value=a); b.cell(row=i,column=2,value=v); b.cell(row=i,column=3,value=c)
box(b,"A4:C6"); box(b,"B4:B6",INP); b["B4"].number_format='"$"#,##0'; b["B5"].number_format="0%"
dv=DataValidation(type="list",formula1='"New agent,Growing agent"',allow_blank=False); b.add_data_validation(dv); dv.add("B6")
hdr(b,8,["Result","Amount",""])
b["A9"]="Marketing budget for the year"; b["B9"]="=B4*B5"
b["A10"]="Marketing budget per month"; b["B10"]="=B9/12"
box(b,"A9:C10"); box(b,"B9:B10",OUT,'"$"#,##0')
hdr(b,12,["Channel","New agent share","Growing agent share","Share for your stage","Monthly amount"])
ch=[("Sphere and referrals (notes, events, gifts)",.30,.20),("Print and branding",.25,.10),("Social media (tools, boosted posts)",.15,.10),
("Website and local SEO",.10,.20),("Direct mail",0,.20),("Paid ads",0,.15),("Test reserve",.20,.05)]
for i,(n,a,c) in enumerate(ch,start=13):
    b.cell(row=i,column=1,value=n); b.cell(row=i,column=2,value=a); b.cell(row=i,column=3,value=c)
    b.cell(row=i,column=4,value='=IF($B$6="New agent",B%d,C%d)'%(i,i)); b.cell(row=i,column=5,value="=D%d*$B$10"%i)
b["A20"]="Total"; b["B20"]="=SUM(B13:B19)"; b["C20"]="=SUM(C13:C19)"; b["D20"]="=SUM(D13:D19)"; b["E20"]="=SUM(E13:E19)"
b["F20"]='=IF(AND(ROUND(B20,4)=1,ROUND(C20,4)=1),"Shares add up to 100%","Shares must add up to 100%")'
box(b,"A13:E20"); box(b,"B13:C19",INP,"0%"); box(b,"D13:D20",OUT,"0%"); box(b,"E13:E20",OUT,'"$"#,##0'); box(b,"B20:C20",OUT,"0%")
for c in b[20]: c.font=Font(bold=True)
b["A22"]="Time is also a budget. A free channel that takes ten hours a week is expensive if you also show homes. Track hours on the Plan sheet."; b["A22"].font=NF
widths(b,{"A":52,"B":22,"C":24,"D":20,"E":16,"F":30})

# 3 Plan
p=wb.create_sheet("3 Plan")
p["A1"]="Your one-page marketing plan"; p["A1"].font=TF
for i,(l,v) in enumerate([("Goal","=\"\"&'1 Goal Backward'!B4&\" closings this year\""),("Audience (one sentence)",""),("Date of first review","")],start=3):
    p.cell(row=i,column=1,value=l).font=Font(bold=True); p.cell(row=i,column=2,value=v)
p.merge_cells("B3:F3"); p.merge_cells("B4:F4"); p.merge_cells("B5:F5")
box(p,"B3:F5",INP)
hdr(p,7,["Channel","Audience","Weekly action","Hours per week","Monthly budget","Result to track"])
ex=[("Sphere and referrals","Past clients and neighbors","Five personal notes or calls",2,100,"Introductions received"),
("Social media","Local first-time buyers","Three posts and ten replies",2,50,"Messages and profile visits"),
("Direct mail","One farm area of 300 homes","Monthly mailing, one drop",1,150,"Calls and valuation requests")]
for i in range(8,18):
    for j in range(6):
        p.cell(row=i,column=j+1,value=ex[i-8][j] if i-8<len(ex) else None)
box(p,"A8:F17",INP,wrap=True); box(p,"D8:D17",None,"0.0"); box(p,"E8:E17",None,'"$"#,##0')
p["A18"]="Total"; p["D18"]="=SUM(D8:D17)"; p["E18"]="=SUM(E8:E17)"
for c in p[18]: c.font=Font(bold=True)
box(p,"A18:F18",OUT); p["D18"].number_format="0.0"; p["E18"].number_format='"$"#,##0'
p["A20"]="Monthly budget from the Budget sheet"; p["E20"]="='2 Budget'!B10"
p["A21"]="Difference (budget minus plan)"; p["E21"]="=E20-E18"
p["F21"]='=IF(E21<0,"Plan is over budget","Plan fits the budget")'
box(p,"A20:F21"); box(p,"E20:E21",OUT,'"$"#,##0')
p["A23"]="Twelve-month view"; p["A23"].font=Font(bold=True,color=NAVY)
hdr(p,24,["Quarter","What happens"])
qs=[("Quarter 1","Set up: clean your contact list, fix your profile, bio, and card, and start the weekly routine."),("Quarter 2","Run the plan and check the numbers at the end of month three."),("Quarter 3","Cut what produced nothing and add one new test with the reserve."),("Quarter 4","Plan next year from the actual numbers, and send your year-end client note.")]
for i,(a,c) in enumerate(qs,start=25):
    p.cell(row=i,column=1,value=a); p.cell(row=i,column=2,value=c); p.merge_cells(start_row=i,start_column=2,end_row=i,end_column=6)
box(p,"A25:F28",None,wrap=True)
widths(p,{"A":26,"B":28,"C":34,"D":16,"E":18,"F":32}); p.freeze_panes="A8"

# 4 Review
r=wb.create_sheet("4 Quarterly Review")
r["A1"]="Review every 90 days"; r["A1"].font=TF
r["A2"]="Record the lead source for every contact, then fill one block per quarter. Change one thing after each review."; r["A2"].font=NF
src=["Sphere and referrals","Social media","Website and local SEO","Paid ads","Direct mail","Print and branding","Other"]
row=4
for q in range(1,5):
    r.cell(row=row,column=1,value="Quarter %d"%q).font=Font(bold=True,color=NAVY,size=12)
    hdr(r,row+1,["Source","Spend","Conversations","Appointments","Closings","Spend per closing","Spend per conversation"])
    s=row+2
    for i,n in enumerate(src):
        rr=s+i; r.cell(row=rr,column=1,value=n)
        r.cell(row=rr,column=6,value='=IF(E%d>0,B%d/E%d,"")'%(rr,rr,rr)); r.cell(row=rr,column=7,value='=IF(C%d>0,B%d/C%d,"")'%(rr,rr,rr))
    t=s+len(src); r.cell(row=t,column=1,value="Total")
    for col in "BCDE": r["%s%d"%(col,t)]="=SUM(%s%d:%s%d)"%(col,s,col,t-1)
    r["F%d"%t]='=IF(E%d>0,B%d/E%d,"")'%(t,t,t); r["G%d"%t]='=IF(C%d>0,B%d/C%d,"")'%(t,t,t)
    box(r,"A%d:G%d"%(s,t)); box(r,"B%d:E%d"%(s,t-1),INP); box(r,"F%d:G%d"%(s,t),OUT)
    for rr in range(s,t+1):
        r["B%d"%rr].number_format='"$"#,##0'; r["F%d"%rr].number_format='"$"#,##0'; r["G%d"%rr].number_format='"$"#,##0.00'
    for c in r[t]: c.font=Font(bold=True)
    r.cell(row=t+1,column=1,value="What I learned, and the one thing I will change:").font=Font(bold=True)
    r.merge_cells(start_row=t+2,start_column=1,end_row=t+2,end_column=7); box(r,"A%d:G%d"%(t+2,t+2),INP)
    r.row_dimensions[t+2].height=34
    row=t+4
widths(r,{"A":28,"B":14,"C":16,"D":16,"E":12,"F":20,"G":22})
for w in wb.worksheets: w.sheet_view.showGridLines=False
wb.save("real-estate-marketing-plan-template.xlsx")
